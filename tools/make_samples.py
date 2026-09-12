"""
สร้างไฟล์ตัวอย่างสำหรับนำเข้า (ข้อมูลสมมติ ไม่ใช่ข้อมูลจริง)
- samples/ip_uc_import_template.xlsx  (ชีตข้อมูล + ชีตคำอธิบาย)
- samples/ip_uc_import_template.csv   (UTF-8 BOM, ตัวเลขมี comma)
- samples/test_invalid_rows.csv        (ไฟล์ทดสอบข้อมูลผิดพลาด)

รัน:  python tools/make_samples.py
"""
import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples"
OUT.mkdir(exist_ok=True)

HEADERS = [
    "งวด STM", "เดือน", "ปีงบประมาณ", "ประเภทบริการ", "BR หลังหักเงินกัน สป.สธ.",
    "BR คูณ K สป.สธ.", "ครั้ง(บริการ)", "Adj.RW ที่ชดเชย", "ชดเชยก่อนหักเงินเดือน",
]

TYPES = [
    ("IP ในเขต", 420, 1.12, 8350, 0.97),
    ("IP ข้ามเขต", 30, 1.60, 9600, 1.00),
    ("ODS", 22, 0.58, 8350, 0.97),
    ("HOMEWARD", 9, 0.95, 8350, 0.97),
    ("NB <= 1,500", 3, 4.40, 8350, 0.97),
    ("NB-HC", 2, 2.10, 8350, 0.97),
    ("UCEP ภาครัฐ", 6, 2.20, 10500, 1.00),
    ("อื่น ๆ", 11, 0.75, 8350, 0.97),
]
MONTHS = [("6810", "2569"), ("6811", "2569"), ("6812", "2569")]

random.seed(69)
rows = []
for code, fy in MONTHS:
    for name, count, cmi, rate, k in TYPES:
        c = max(1, round(count * random.uniform(0.9, 1.1)))
        adj = round(c * cmi * random.uniform(0.93, 1.07), 4)
        br_after = round(adj * rate * 0.92, 2)
        br_k = round(br_after * k, 2)
        comp = round(br_k * random.uniform(1.0, 1.03), 2)
        rows.append([f"{code}_IP_01", code, fy, name, br_after, br_k, c, adj, comp])

# ------------------------------ Excel ------------------------------
NAVY = "0F2747"
thin = Side(style="thin", color="E2E8F0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "ข้อมูลนำเข้า"
ws.append(HEADERS)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
ws.row_dimensions[1].height = 34

for r in rows:
    ws.append(r)
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for i, cell in enumerate(row):
        cell.border = border
        if i <= 3:
            cell.number_format = "@"
        elif i == 6:
            cell.number_format = "#,##0"
        elif i == 7:
            cell.number_format = "#,##0.0000"
        else:
            cell.number_format = "#,##0.00"

for col, width in zip("ABCDEFGHI", [16, 9, 12, 18, 24, 20, 13, 17, 24]):
    ws.column_dimensions[col].width = width
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

guide = wb.create_sheet("คำอธิบาย")
guide.append(["คำอธิบายไฟล์นำเข้า IP.UC Performance Analytics Dashboard"])
guide["A1"].font = Font(bold=True, size=14, color=NAVY)
guide.append(["ข้อมูลในชีต “ข้อมูลนำเข้า” เป็นตัวเลขสมมติเพื่อแสดงรูปแบบเท่านั้น กรุณาลบแล้ววางข้อมูลจริงจากรายงาน STM"])
guide.append([])
guide.append(["หัวคอลัมน์ (ต้องตรงตามนี้)", "ฟิลด์ในระบบ", "จำเป็น", "รูปแบบ / ตัวอย่าง"])
for cell in guide[4]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
spec = [
    ("งวด STM", "stm_period", "ใช่", "ข้อความ เช่น 6907_IP_02"),
    ("เดือน", "month_code", "ใช่", "YYMM (พ.ศ. 2 หลัก) เช่น 6907 = ก.ค. 2569 · เว้นว่างได้ถ้างวด STM ขึ้นต้นด้วย YYMM"),
    ("ปีงบประมาณ", "fiscal_year", "ใช่", "พ.ศ. 4 หลัก เช่น 2569 (ต.ค. 2568 – ก.ย. 2569) · เว้นว่างได้ ระบบคำนวณจากเดือน"),
    ("ประเภทบริการ", "service_type", "ใช่ (หัวคอลัมน์)", "IP ในเขต, IP ข้ามเขต, ODS, HOMEWARD, NB <= 1,500, NB-HC, UCEP ภาครัฐ, อื่น ๆ · ค่าว่าง = ไม่ระบุประเภทบริการ"),
    ("BR หลังหักเงินกัน สป.สธ.", "br_after_deduction", "ใช่", "ตัวเลข (บาท) รองรับ 1,184,501.66"),
    ("BR คูณ K สป.สธ.", "br_k", "ใช่", "ตัวเลข (บาท)"),
    ("ครั้ง(บริการ)", "service_count", "ใช่", "จำนวนครั้งบริการ"),
    ("Adj.RW ที่ชดเชย", "adj_rw", "ใช่", "ตัวเลขทศนิยมสูงสุด 4 ตำแหน่ง"),
    ("ชดเชยก่อนหักเงินเดือน", "compensation", "ใช่", "ตัวเลข (บาท)"),
]
for s in spec:
    guide.append(list(s))
guide.append([])
guide.append(["กติกาการตรวจสอบ"])
guide.cell(row=guide.max_row, column=1).font = Font(bold=True, color=NAVY)
for line in [
    "คีย์ข้อมูลซ้ำ = งวด STM + ประเภทบริการ (ไม่สนตัวพิมพ์เล็ก/ใหญ่)",
    "ช่องตัวเลขที่ว่างหรือเป็น “-” จะถือเป็น 0 · ข้อความที่ไม่ใช่ตัวเลขจะไม่ผ่านการตรวจสอบ",
    "แถว “รวม” / “Total” จะไม่ถูกนำเข้า · แถวว่างจะถูกข้าม",
    "Append = เพิ่มเฉพาะข้อมูลใหม่ · Upsert = เพิ่มใหม่หรืออัปเดตข้อมูลที่ซ้ำ · Replace Batch = แทนที่ข้อมูลทั้งหมดของงวด STM ที่เลือก",
]:
    guide.append(["• " + line])
guide.column_dimensions["A"].width = 30
guide.column_dimensions["B"].width = 22
guide.column_dimensions["C"].width = 16
guide.column_dimensions["D"].width = 90

wb.save(OUT / "ip_uc_import_template.xlsx")

# ------------------------------ CSV ------------------------------
def money(v):
    return f"{v:,.2f}"

with open(OUT / "ip_uc_import_template.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
    w.writerow(HEADERS)
    for r in rows:
        w.writerow([r[0], r[1], r[2], r[3], money(r[4]), money(r[5]), f"{r[6]:,}", f"{r[7]:.4f}", money(r[8])])

# ------------------------------ ไฟล์ทดสอบข้อมูลผิดพลาด ------------------------------
with open(OUT / "test_invalid_rows.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["รายงาน STM ผู้ป่วยใน (ไฟล์ทดสอบข้อมูลผิดพลาด — ห้ามใช้เป็นข้อมูลจริง)"])
    w.writerow(HEADERS)
    w.writerow(["6901_IP_01", "6901", "2569", "IP ในเขต", "3,450,120.50", "3,346,616.89", "412", "455.1234", "3,380,083.06"])   # ผ่าน
    w.writerow(["6901_IP_01", "6901", "2569", "IP ในเขต", "1", "1", "1", "1", "1"])                                          # ซ้ำในไฟล์
    w.writerow(["6901_IP_01", "6913", "2569", "ODS", "10", "10", "1", "0.5", "10"])                                          # เดือนผิด
    w.writerow(["", "6901", "2569", "HOMEWARD", "10", "10", "1", "0.5", "10"])                                               # ไม่มีงวด STM
    w.writerow(["6901_IP_01", "6901", "2569", "NB-HC", "abc", "10", "1", "0.5", "10"])                                       # ไม่ใช่ตัวเลข
    w.writerow(["6901_IP_01", "6901", "2568", "UCEP ภาครัฐ", "25,000", "25,000", "3", "2.4", "25,500"])                      # คำเตือน: ปีงบไม่สอดคล้อง
    w.writerow(["6901_IP_01", "6901", "2569", "", "5,000", "5,000", "2", "0.9", "5,100"])                                    # คำเตือน: ไม่ระบุประเภท
    w.writerow(["6901_IP_01", "6901", "2569", "อื่น ๆ", "(1,234.00)", "-", "1", "0.2", "0"])                                 # คำเตือน: ติดลบ
    w.writerow([])                                                                                                             # แถวว่าง
    w.writerow(["รวม", "", "", "รวม", "3,480,000", "3,380,000", "420", "460", "3,410,000"])                                  # แถวรวม

print("created:", *(p.name for p in sorted(OUT.iterdir())))
